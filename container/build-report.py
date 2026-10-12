"""Build the Lab 3 report from preserved native and actual Docker measurements."""
import argparse,json,subprocess,sys
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
NAMES=('smoke','load','stress','spike','endurance')
def read(path,default):return json.loads(path.read_text(encoding='utf-8-sig')) if path.exists() else default
def fmt(value,percent=False,sign=False):
    if value is None:return 'Не виміряно'
    value=value*100 if percent else value
    return f'{value:+.2f}' if sign else f'{value:.2f}'
def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for index,row in enumerate(t.rows):
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margin.append(el)
            props.append(margin)
            if index==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7E6E6');props.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.size=Pt(9);run.bold=index==0
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    gap=doc.add_paragraph();gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.space_after=Pt(4)
    return t

def hyperlink(paragraph,text,url):
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color);run.append(props)
    value=OxmlElement('w:t');value.text=text;run.append(value);el.append(run);paragraph._p.append(el)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'TravelerAPI_Lab3_Report.docx');a=parser.parse_args()
    subprocess.run([sys.executable,str(ROOT/'container/compare-results.py')],check=True)
    comparison=read(ROOT/'results/comparison.json',{});rows=comparison['rows'];bmeta=comparison['baseline_meta'];dmeta=comparison['docker_meta']
    bdata={r['test']:r for r in read(ROOT/'results/baseline/analysis.json',[])}
    ddata={r['test']:r for r in read(ROOT/'results/docker/analysis.json',[])}
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.keep_together=True
    for name,size in [('Title',18),('Heading 1',13),('Heading 2',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        for border in list(style._element.get_or_add_pPr().findall(qn('w:pBdr'))):style._element.get_or_add_pPr().remove(border)
    doc.core_properties.author='Ольга Великодна';doc.core_properties.title='Лабораторна робота 3 Контейнеризація TravelerAPI'
    def p(text):return doc.add_paragraph(text)
    def h(text):doc.add_heading(text,1)
    def code(text):
        paragraph=p(text)
        for run in paragraph.runs:run.font.name='Consolas';run.font.size=Pt(9)
    doc.add_paragraph('Лабораторна робота 3',style='Title')
    p('Контейнеризація TravelerAPI та порівняння продуктивності')
    p('Виконала Ольга Великодна')
    h('Мета та стан виконання')
    p('Мета роботи — створити Dockerfile і повний Docker Compose стек API та PostgreSQL, перевірити його працездатність і оцінити зміну продуктивності відносно нативної лабораторної 2. Підготовлено конфігурацію контейнерів, моніторинг, повторний запуск тих самих k6-сценаріїв та автоматичне порівняння.')
    ready=[r['test'] for r in rows if r['comparison_status']=='ready']
    p('Порівняння з фактичними Docker-результатами доступне для: '+(', '.join(ready) if ready else 'жодного сценарію. Нативні дані вже збережено; контейнерні вимірювання та скриншоти очікують локального запуску. Відсоток впливу Docker до цього запуску не встановлено.'))
    h('Dockerfile для додатку')
    p('Dockerfile має два етапи. dependencies створює /opt/venv і встановлює закріплені залежності requirements.lock.txt. runtime отримує готове середовище, app.py та SQL/OpenAPI-файли. Базовий Python 3.12.10 відповідає нативному baseline. Застосунок запускається користувачем UID 1000, а не root, через Uvicorn на 4567 з одним worker. Число з’єднань пулу залишається 20. HTTP healthcheck використовує стандартну бібліотеку Python.')
    h('Docker Compose для повного стеку')
    p('docker-compose.yml визначає web і postgres у спільній bridge-мережі. PostgreSQL 18.4 має named volume pgdata, змінні з .env і healthcheck pg_isready. Його порт не публікується у Windows. Для PostgreSQL 18 том змонтовано в /var/lib/postgresql, де образ зберігає версійну папку даних. web починає роботу після service_healthy БД і публікує лише 127.0.0.1:4567.')
    p('API має read_only, tmpfs /tmp, no-new-privileges та ротацію журналів. Базова конфігурація обмежує web двома CPU і 1 ГБ RAM, postgres — двома CPU і 2 ГБ RAM. benchmark-override прибирає ці per-container caps для зіставлення з нативними процесами; загальні ресурси Docker Desktop залишаються частиною умов експерименту. Окремий dev-override додає bind mounts і --reload та не застосовується для вимірювань.')
    p('.env створюється з випадковим паролем, не публікується і виключений із build context. .env.example містить усі ключі. Приклад не містить справжніх облікових даних користувача; URL і паролі не записані літералами у Compose.')
    doc.add_page_break()
    h('Запуск перевірка та моніторинг')
    code('powershell -ExecutionPolicy Bypass -File .\\container\\setup-env.ps1')
    code('powershell -ExecutionPolicy Bypass -File .\\container\\start-stack.ps1')
    p('Скрипт запускає benchmark-профіль через docker compose up --build -d --wait. Для базового обмеженого стеку достатньо створити .env і виконати docker compose up --build. Адреса перевірки — http://localhost:4567/health, очікувана відповідь status=ok.')
    code('powershell -ExecutionPolicy Bypass -File .\\container\\check-stack.ps1')
    p('Перевірка Hurl виконує ті самі CRUD і validation тести лабораторної 1 з виправленим API та додатковий ordering.hurl. Справжні одночасні запити перевіряє окремий test_concurrency.py; послідовний race-conditions.hurl не називається доказом одночасності.')
    exitfile=ROOT/'results/monitoring/hurl-exit-code.txt'
    if exitfile.exists():p('Зафіксований код завершення контейнерного Hurl: '+exitfile.read_text(encoding='utf-8-sig').strip()+'. Повний протокол — results/monitoring/hurl-console.txt.')
    else:p('Контейнерний запуск Hurl ще не надано. Історичні успішні перевірки Lab 2 не замінюють перевірку контейнерів; скриншот успішного Hurl не заявляється.')
    code('docker compose ps\ndocker compose logs -f web\ndocker stats $(docker compose ps -q)')
    p('Скрипт collect-monitoring.ps1 зберігає compose ps, журнали web і docker stats у results/monitoring. run-performance.ps1 збирає JSON-знімки stats кожні приблизно 5 секунд. Скриншоти додаються до звіту з compose-ps.png, web-logs.png, docker-stats.png та hurl.png. Ctrl+C припиняє перегляд журналу або stats, а не сам стек.')
    h('Коміти лабораторної роботи')
    p('Вихідний репозиторій: https://github.com/Olyalya613/RSD_lab_1. Baseline отримано з опублікованої гілки lab2-corrected, коміт 4fa48edf2e5dfdcf78325acde264215e86164117. Нові зміни підготовлено в локальній гілці lab3-docker; їхня віддалена публікація ще не виконана.')
    try:
        lines=subprocess.check_output(['git','log','--format=%H|%s','origin/lab2-corrected..HEAD'],cwd=ROOT,text=True).splitlines()
        commits=[line.split('|',1) for line in lines if any(title in line for title in ['Add non-root Dockerfile','Add container Hurl','Preserve actual Lab 2 baseline','Add Docker measurements'])]
    except (OSError,subprocess.SubprocessError):commits=[]
    t=table(doc,['Коміт','Зміна'],[(sha[:7],message) for sha,message in commits] or [('Очікується','Локальна історія буде збережена у .git та history.bundle')],[2.8,14.2])
    for i,(sha,message) in enumerate(commits,1):
        paragraph=t.rows[i].cells[0].paragraphs[0];paragraph.clear();hyperlink(paragraph,sha[:7],'https://github.com/Olyalya613/RSD_lab_1/commit/'+sha)
    p('Посилання на нові коміти запрацюють після git push -u origin lab3-docker. Очікувана адреса гілки: https://github.com/Olyalya613/RSD_lab_1/tree/lab3-docker. Репозиторій має app.py і Compose-файли безпосередньо в корені, без зовнішньої папки.')
    doc.add_page_break()
    h('Середовище та методика порівняння')
    table(doc,['Параметр','Нативний Lab 2','Docker'],[(label,str(bmeta.get(key,'Не зафіксовано')),str(dmeta.get(key,'Очікується запуск'))) for label,key in [('CPU','cpu'),('Логічні процесори','logical_processors'),('RAM у ГБ','ram_gb'),('Python','python_version'),('PostgreSQL','postgresql_version'),('Ізоляція','isolation'),('k6','k6_version'),('Workers','workers'),('Пул БД','db_pool_max')]],[3.3,6.85,6.85])
    if dmeta:p('Ресурси Docker Engine: '+str(dmeta.get('docker_cpus','невідомо'))+' CPU, '+str(dmeta.get('docker_ram_gb','невідомо'))+' ГБ RAM. Профіль: '+str(dmeta.get('profile','невідомо'))+'. Ідентифікатори образів і фактичні ліміти записані в *-environment.json.')
    p('k6 запускається у Windows на тому самому комп’ютері; API та БД перенесено в Linux-контейнери Docker Desktop. app.py, схема БД, версії залежностей, один worker, пул 20, профілі VU і think time 0,5 с збережені. Використано ті самі Smoke, Load, Stress, Spike та 30-хвилинний Endurance. Setup/teardown виключені з основних метрик через measured:true.')
    p('Порівнюються avg і p95 окремо: p95 — це 95-й перцентиль, а не середнє. Зміна часу у % обчислюється як (Docker − native) / native × 100. Частка помилок порівнюється у процентних пунктах. Позитивна зміна часу означає уповільнення, від’ємна — прискорення. Правильні 409 конфлікти версій рахуються окремо від неочікуваних HTTP-помилок.')
    p('Результат характеризує весь перехід Windows native → Docker Desktop Linux, включно з ОС, віртуалізацією, мережею та сховищем. Один запуск не ізолює чисту накладну вартість контейнера. Базовий вміст БД, стан кешів, locale та фонові процеси також впливають; нативна metadata не містить locale і per-test хешів. Після цього експерименту варто повторити обидві серії за однакових початкових умов, якщо потрібна оцінка розкиду.')
    doc.add_page_break()
    h('Показники результатів тестування')
    table(doc,['Тест','Native p95 мс','Docker p95 мс','Зміна p95 %'],[(r['test'],fmt(r['baseline'].get('p95_ms')),fmt(r['docker'].get('p95_ms')),fmt(r['delta'].get('p95_ms_percent'),sign=True)) for r in rows],[3,4,4,6])
    table(doc,['Тест','Native avg мс','Docker avg мс','Зміна avg %'],[(r['test'],fmt(r['baseline'].get('avg_ms')),fmt(r['docker'].get('avg_ms')),fmt(r['delta'].get('avg_ms_percent'),sign=True)) for r in rows],[3,4,4,6])
    table(doc,['Тест','Native помилки %','Docker помилки %','Зміна п п'],[(r['test'],fmt(r['baseline'].get('error_rate'),percent=True),fmt(r['docker'].get('error_rate'),percent=True),fmt(r['delta'].get('errors_percentage_points'),sign=True)) for r in rows],[3,4,4,6])
    for r in rows:
        if r['comparison_status']=='different_conditions':p(r['test']+': різні умови; '+', '.join(r['problems'])+'. Відсотковий висновок не сформовано.')
    p('Наведені нативні числа є результатами користувача з analysis.json, а не вимірюваннями середовища підготовки. У нативному Spike зафіксовано '+fmt(bdata.get('spike',{}).get('overall',{}).get('error_rate'),percent=True)+'% невдалих запитів. Тому baseline не можна описувати як стабільний для всіх профілів. Без контейнерних даних не встановлено, чи змінилася ця частка після перенесення.')
    h('Перевірка підготовленого комплекту')
    p('Docker Compose v2.39.0 прийняв базовий, benchmark і dev варіанти командою config --quiet. Пройшли 17 Python-перевірок валідації та аналізу, з них 8 нових перевірок порівняння. 11 інтеграційних тестів PostgreSQL у цій підготовці пропущено через відсутність запущеного тестового API/БД. Перевірено синтаксис Python і PowerShell. Docker Engine недоступний: збірка образу й фактичний запуск контейнерів тут не виконувалися.')
    h('Висновки')
    if ready:
        for r in rows:
            if r['comparison_status']=='ready':p(r['test']+': зміна p95 '+fmt(r['delta'].get('p95_ms_percent'),sign=True)+'%, зміна середнього часу '+fmt(r['delta'].get('avg_ms_percent'),sign=True)+'%, зміна помилок '+fmt(r['delta'].get('errors_percentage_points'),sign=True)+' п. п. Це результат даного запуску за зафіксованих умов.')
    else:p('Підготовлено відтворюваний Docker-стек і інструмент порівняння. Нативний baseline збережено без підміни результатів. Остаточний висновок про зміну швидкодії буде сформовано після запуску того самого навантаження на контейнерах і додавання фактичних метрик та скриншотів.')
    h('Графіки та скриншоти')
    inserted=False
    for name in NAMES:
        before=bdata.get(name,{});after=ddata.get(name,{})
        if len(before.get('timeline',[]))>1 or len(after.get('timeline',[]))>1:
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt
            fig,axes=plt.subplots(2,1,figsize=(7,3.5),sharex=True)
            for label,data in [('Native Lab 2',before),('Docker Lab 3',after)]:
                timeline=data.get('timeline',[])
                if not timeline:continue
                origin=timeline[0]['time'];x=[(point['time']-origin)/60 for point in timeline]
                axes[0].plot(x,[point['p95_ms'] for point in timeline],label=label)
                axes[1].plot(x,[100*(point['error_rate'] or 0) for point in timeline],label=label)
            axes[0].set_ylabel('p95 (ms)');axes[0].set_ylim(bottom=0);axes[1].set_ylabel('Errors (%)');axes[1].set_ylim(bottom=0);axes[1].set_xlabel('Minutes');axes[0].legend();axes[0].grid(alpha=.25);axes[1].grid(alpha=.25);fig.suptitle(name.capitalize());fig.tight_layout()
            image=ROOT/'results'/f'comparison-{name}.png';fig.savefig(image,dpi=150);plt.close(fig)
            caption=p('Графік '+name+' з фактичного 30-секундного ряду метрик. Якщо Docker-серії немає, показано тільки нативний baseline.');caption.paragraph_format.keep_with_next=True
            doc.add_picture(str(image),width=Cm(16));inserted=True
    for name in ['compose-ps','web-logs','docker-stats','hurl']:
        image=ROOT/'results/monitoring'/f'{name}.png'
        if image.exists():
            caption=p('Скриншот контейнерного запуску '+name);caption.paragraph_format.keep_with_next=True;doc.add_picture(str(image),width=Cm(16));inserted=True
    for name in NAMES:
        image=ROOT/'results/docker'/f'{name}-screenshot.png'
        if image.exists():
            caption=p('Скриншот k6 у контейнерному експерименті '+name);caption.paragraph_format.keep_with_next=True;doc.add_picture(str(image),width=Cm(16));inserted=True
    if not any((ROOT/'results/monitoring'/f'{name}.png').exists() for name in ['compose-ps','web-logs','docker-stats','hurl']):p('Обов’язкові скриншоти Docker-моніторингу та Hurl ще не надані. Після збереження PNG у results/monitoring і повторного build-report.py вони автоматично увійдуть у документ.')
    h('Джерела')
    for source in ['Методичні вказівки SumDU.CS.DDS.Lab3.pdf','https://github.com/Olyalya613/RSD_lab_1/tree/lab2-corrected','https://docs.docker.com/build/building/multi-stage/','https://docs.docker.com/compose/how-tos/startup-order/','https://hub.docker.com/_/postgres','https://docs.docker.com/desktop/features/vmm/']:
        paragraph=p(source);paragraph.paragraph_format.keep_with_next=source!='https://docs.docker.com/desktop/features/vmm/'
    a.output.parent.mkdir(parents=True,exist_ok=True);doc.save(a.output);print(a.output)
if __name__=='__main__':main()
