"""Generate a Word report from actual local results, with unmeasured fields explicit."""
import argparse
import json
import subprocess
import sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
PERF=ROOT/'tests/performance-tests'
NAMES=('smoke','load','stress','spike','endurance')

def number(value,unit='',percent=False):
    if value is None:return 'Не виміряно'
    if isinstance(value,bool):return 'Так' if value else 'Ні'
    return f'{value*100 if percent else value:.2f}{unit}'

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,label in zip(t.rows[0].cells,headers):c.text=str(label)
    for row in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,row):cell.text=str(value)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
    props.append(borders)
    for index,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cellprops=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'90');margin.set(qn('w:type'),'dxa');margins.append(margin)
            cellprops.append(margins)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cellprops.append(shade)
                for run in cell.paragraphs[0].runs:run.bold=True
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:run.font.size=Pt(9)
    return t

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--results',type=Path,default=PERF/'results');parser.add_argument('--output',type=Path,default=ROOT/'TravelerAPI_Lab2_Report.docx');args=parser.parse_args()
    subprocess.run([sys.executable,str(ROOT/'native/analyze-results.py'),'--results',str(args.results)],check=True)
    results={r['test']:r for r in json.loads((args.results/'analysis.json').read_text(encoding='utf-8'))}
    meta_file=args.results/'run-meta.json';meta=json.loads(meta_file.read_text(encoding='utf-8-sig')) if meta_file.exists() else {}
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10.5);style.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',18),('Heading 1',13),('Heading 2',11)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        for border in list(st._element.get_or_add_pPr().findall(qn('w:pBdr'))):st._element.get_or_add_pPr().remove(border)
    def p(text):doc.add_paragraph(text)
    def h(text):doc.add_heading(text,1)
    def code(text):
        x=doc.add_paragraph();r=x.add_run(text);r.font.name='Consolas';r.font.size=Pt(9)
    doc.add_paragraph('Лабораторна робота 2',style='Title')
    p('Тестування продуктивності TravelerAPI за допомогою k6')
    h('Мета та стан виконання')
    p('Мета роботи — підготувати та виміряти baseline продуктивності API для наступного порівняння архітектур. Реалізовано Smoke, Load, Stress, Spike та 30-хвилинний Endurance сценарії, експорт метрик і побудову звіту. API, PostgreSQL та k6 потрібно запускати як звичайні процеси на одному комп’ютері Windows без Docker, WSL чи віртуальних машин.')
    measured=[n for n in NAMES if results[n]['status']=='measured']
    p('Фактичні вимірювання наявні для: '+(', '.join(measured) if measured else 'жодного сценарію. Значення нижче заповняться після локального запуску; результати не підмінено припущеннями.'))
    h('Середовище та методика')
    table(doc,['Параметр','Значення'],[(label,meta.get(key,'Фіксується перед запуском')) for label,key in [('Операційна система','os'),('Процесор','cpu'),('Логічні процесори','logical_processors'),('RAM у ГБ','ram_gb'),('PostgreSQL','postgresql_version'),('Python','python_version'),('k6','k6_version'),('Ізоляція транзакцій','isolation')]])
    p('API слухає 127.0.0.1:4567, PostgreSQL — 127.0.0.1:5432. Використано один Uvicorn worker і пул до 20 з’єднань. Тип навантаження — closed model ramping-vus; 70% ітерацій читають плани, 20% створюють та редагують власний план, 10% редагують спільний план і локацію. Між ітераціями є пауза 0,5 с. Запущений на тому самому комп’ютері k6 також використовує його ресурси.')
    table(doc,['Сценарій','Профіль за замовчуванням'],[
      ('Smoke','2 VU, 20 секунд, CRUD та конфлікти версій'),
      ('Load','10 → 50 → 100 VU з утриманням кожного рівня, 4 хв 30 с'),
      ('Stress','100 → 200 → 400 → 600 → 1000 VU, потім відновлення на 10 VU'),
      ('Spike','10 → 1000 VU за 2 с, сплеск 30 с, відновлення 60 с'),
      ('Endurance','1 хв розігріву + 28 хв на 100 VU + 1 хв зниження = 30 хв')])
    doc.add_page_break()
    h('Завдання 1 Smoke та адаптація API')
    p('Конфігурацію адаптовано до порту 4567 і /health зі значенням status=ok. При оновленні локації передається locations.version. Smoke окремо перевіряє застарілу версію плану й локації, очікуючи 409, а після видалення — 404. Для цих запитів налаштовано expectedStatuses, тому правильні 409 та 404 не збільшують http_req_failed. Конфлікти обліковуються окремо, а їхній Rate має знаменник з усіх PUT-запитів.')
    code('k6 run --no-usage-report --out json=results/smoke-metrics.json.gz smoke-test.js')
    table(doc,['Тест','p95 у мс','Помилки','Статус'],[(n,number(results[n].get('overall',{}).get('p95_ms')),number(results[n].get('overall',{}).get('error_rate'),'%',True),'Виміряно' if results[n]['status']=='measured' else 'Неповний запуск' if results[n]['status']=='incomplete' else 'Не виміряно') for n in NAMES])
    h('Завдання 2 Load Testing')
    load=results['load'];peak=load.get('peak') or {}
    table(doc,['Показник','Результат'],[
      ('Найбільший перевірений рівень VU в межах SLO',number(load.get('largest_tested_vus_within_slo')) if load.get('largest_tested_vus_within_slo') is not None or load['status']!='measured' else 'Не встановлено за критеріями SLO'),
      ('Середній час відповіді на піку',number(peak.get('avg_ms'),' мс')),
      ('p95 на піку',number(peak.get('p95_ms'),' мс')),
      ('Частка помилок на піку',number(peak.get('error_rate'),'%',True))])
    p('Рівень вважається прийнятним, якщо p95 < 200 мс, частка помилок < 1%, отримано не менш ніж 20 запитів і зафіксовано досягнення заданої кількості VU. Це найбільший прийнятний рівень серед перевірених, а не доведена абсолютна межа API.')
    h('Завдання 3 Stress Testing')
    stress=results['stress']
    table(doc,['Показник','Результат'],[
      ('Перший перевірений рівень деградації',number(stress.get('degradation_vus')) if stress.get('degradation_vus') is not None or stress['status']!='measured' else 'Не виявлено в отриманих даних'),
      ('Операційна точка відмови',number(stress.get('operational_breaking_point_vus')) if stress.get('operational_breaking_point_vus') is not None or stress['status']!='measured' else 'Не досягнуто в отриманих даних'),
      ('Відновлення в межах SLO після зниження',number(stress.get('recovered_within_slo')))])
    p('Деградація — перше порушення SLO на утримуваному рівні. Для цього досліду операційну точку відмови визначено як ≥5% помилок або p95 ≥2000 мс на утримуваному рівні. Якщо значення не досягнуто, межу слід позначити як не встановлену в перевіреному діапазоні, а не прирівнювати до найбільшого VU.')
    doc.add_page_break()
    h('Завдання 4 Spike Testing')
    spike=results['spike'];sr=spike.get('spike') or {}
    table(doc,['Показник','Результат'],[
      ('p95 під час сплеску',number(sr.get('p95_ms'),' мс')),
      ('Частка помилок під час сплеску',number(sr.get('error_rate'),'%',True)),
      ('Час підтвердженого відновлення',number(spike.get('recovery_seconds'),' с') if spike.get('recovery_seconds') is not None or spike['status']!='measured' else 'Не зафіксовано за час спостереження')])
    p('Час відновлення вимірюється після повернення на 10 VU. Потрібні три послідовні 5-секундні вікна з не менш ніж 5 запитами кожне, часткою помилок <1% і p95 не вище max(200 мс, 1,2 × baseline p95). Тайм-аути й неочікувані відповіді аналізуються як невдалі запити. Ці дані самі по собі не доводять втрату пакетів у мережі.')
    h('Завдання 5 Endurance Testing')
    endurance=results['endurance'];er=endurance.get('phases',{}).get('steady',{})
    table(doc,['Показник','Результат'],[
      ('p95 на початку стабільної фази',number(endurance.get('first_five_minutes_p95_ms'),' мс')),
      ('p95 наприкінці стабільної фази',number(endurance.get('last_five_minutes_p95_ms'),' мс')),
      ('Співвідношення останнього p95 до першого',number(endurance.get('p95_growth_ratio'))),
      ('Частка помилок стабільної фази',number(er.get('error_rate'),'%',True))])
    p('Для початкового й кінцевого порівняння використовуються повні 30-секундні вікна всередині відповідних 5-хвилинних інтервалів. Зростання p95 більш ніж на 20% — сигнал можливої деградації, який слід зіставити з CPU, RAM і часткою помилок. Моніторинг збирає сумарну пам’ять процесів Python і PostgreSQL та загальне навантаження CPU кожні 5 секунд. Зростання пам’яті без додаткової діагностики не підтверджує витік.')
    h('Графіки та скриншоти фактичних запусків')
    inserted=False
    for name in NAMES:
        r=results[name]
        if r['status']=='measured' and r.get('timeline'):
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt
            ts=r['timeline'];origin=ts[0]['time'];x=[(t['time']-origin)/60 for t in ts]
            fig,axes=plt.subplots(2,1,figsize=(7,4),sharex=True)
            axes[0].plot(x,[t['p95_ms'] for t in ts]);axes[0].set_ylabel('p95 (ms)');axes[0].grid(alpha=.25)
            axes[1].plot(x,[100*(t['error_rate'] or 0) for t in ts]);axes[1].set_ylabel('Errors (%)');axes[1].set_xlabel('Minutes');axes[1].grid(alpha=.25)
            fig.suptitle(name.capitalize());fig.tight_layout();image=args.results/f'{name}-chart.png';fig.savefig(image,dpi=180);plt.close(fig)
            doc.add_picture(str(image),width=Cm(16));inserted=True
        shot=args.results/f'{name}-screenshot.png'
        if shot.exists():p(f'Скриншот запуску {name}');doc.add_picture(str(shot),width=Cm(16));inserted=True
    if not inserted:p('Графіки додаються з реального JSON-потоку після запуску. Для включення скриншотів потрібно зберегти зображення як smoke-screenshot.png, load-screenshot.png, stress-screenshot.png, spike-screenshot.png та endurance-screenshot.png у папці results і повторно згенерувати звіт.')
    doc.add_page_break()
    h('Структура проєкту та коміти')
    p('app.py і docs/schema.sql — TravelerAPI та PostgreSQL. native — ініціалізація БД, запуск API, запис середовища, аналіз і генератор звіту. tests/performance-tests/config — адреси та профілі; utils — клієнт API, генератор даних, користувацький сценарій; *-test.js — п’ять основних тестів; results — фактичні метрики, журнали й скриншоти. Зайву вкладеність репозиторію усунено: ці файли розміщені безпосередньо в корені. Dockerfile і compose.yaml не входять до цього комплекту.')
    try:
        log=subprocess.check_output(['git','log','--format=%H|%s'],cwd=ROOT,text=True)
        patterns=['Fix TravelerAPI CAS', 'add k6 performance testing baseline','Add Load testing','Add Stress testing','Add Spike testing','Add Endurance testing']
        commits=[line.split('|',1) for line in log.splitlines() if any(p in line for p in patterns)]
    except (OSError,subprocess.SubprocessError):commits=[]
    commit_table=table(doc,['Коміт','Зміна'],[(item[0][:7],item[1]) for item in commits] or [('Перевірити git log','Історія міститься в .git або history.bundle')])
    for index,(sha,message) in enumerate(commits,1):
        paragraph=commit_table.rows[index].cells[0].paragraphs[0];paragraph.clear()
        link=OxmlElement('w:hyperlink')
        relation=paragraph.part.relate_to('https://github.com/Olyalya613/RSD_lab_1/commit/'+sha,RT.HYPERLINK,is_external=True)
        link.set(qn('r:id'),relation)
        run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color);run.append(props)
        text=OxmlElement('w:t');text.text=sha[:7];run.append(text);link.append(run);paragraph._p.append(link)
    p('Хеші в таблиці є посиланнями на очікувані адреси GitHub. Вони запрацюють лише після публікації цієї локальної історії; на момент підготовки віддалені коміти не створені.')
    p('Робота підготовлена в локальній гілці lab2-corrected. Віддалена публікація не виконувалася. Після git push -u origin lab2-corrected репозиторій буде доступний за адресою https://github.com/Olyalya613/RSD_lab_1/tree/lab2-corrected. Адреси комітів формуються як https://github.com/Olyalya613/RSD_lab_1/commit/ІДЕНТИФІКАТОР і перевіряються після публікації.')
    h('Виправлення конкурентного доступу')
    p('Обробники, що використовують синхронний psycopg, оголошені def. FastAPI виконує їх у thread pool, тому один Uvicorn worker обслуговує одночасні запити. Асинхронна залежність payload лише читає JSON і не працює з БД.')
    p('Звичайне редагування локації виконується через UPDATE locations SET ..., version=version+1 WHERE id=%s AND version=%s RETURNING *. SELECT FOR UPDATE на плані чи локації в цій гілці відсутній. Якщо UPDATE не повернув рядок, API повертає 409 із current_version, або 404, якщо локації вже немає.')
    p('GET плану використовує один SELECT із LATERAL та jsonb_agg. Поля плану, order_version і впорядковані локації читаються з одного MVCC-знімка за READ COMMITTED; двох окремих знімків більше немає.')
    p('Порядок списку захищає travel_plans.order_version. PUT із visit_order вимагає version локації й order_version колекції. Спочатку CAS резервує колекцію, потім CAS оновлює переміщувану локацію; сусіди зсуваються й отримують нові версії. Усе відбувається в одній транзакції. Після будь-якої помилки зміни обох CAS відкочуються. Додавання й видалення також збільшують order_version. Песимістичний замок батьківського плану потрібен для структурних операцій, а не для зміни notes.')
    h('Перевірка підготовленого коду')
    p('Протокол перевірки міститься в docs/verification-unittest.txt та docs/verification-k6-*.txt. Ці журнали є перевіркою правильності, а не вимірюванням baseline на комп’ютері користувача. Точні умови стенда описано в docs/VERIFICATION.md. Для повторення конкурентних перевірок наведено команди в START_HERE_UA.md.')
    p('Перевірено 20 Python-тестів: 5 перевірок валідації, 4 перевірки аналізатора та 11 інтеграційних перевірок зі справжніми HTTP-запитами й PostgreSQL. Серед них — одночасні CAS однієї локації, конфлікт різних перестановок з одним токеном колекції, робота інших запитів під час очікування замка, узгоджений знімок при паралельних commit, ущільнення порядку після видалення та rollback.')
    p('Усі п’ять Hurl-файлів також пройшли: 69 HTTP-запитів, 0 невдалих файлів. У короткому VERIFY-режимі всі п’ять k6-сценаріїв виконали 100% перевірок без неочікуваних відповідей. Smoke додатково перевіряє відмову при застарілому order_version іншої локації.')
    p('Короткий режим VERIFY=1 призначений лише для перевірки сценаріїв та виключається аналізатором із baseline. Реальні Load, Stress, Spike і 30-хвилинний Endurance потрібно виконати на нативному Windows-середовищі. Значення швидкодії до цього запуску не заявляються.')
    h('Висновок')
    p('Підготовлено відтворюваний інструмент вимірювання продуктивності TravelerAPI. Числові висновки щодо кількості користувачів, точки відмови та стабільності формуються з фактичних вимірювань у таблицях цього звіту. До запуску на нативному середовищі такі висновки не встановлено.')
    h('Джерела')
    for link in ['https://github.com/potapuff/SumDU.DDS.25F.lab2','https://grafana.com/docs/k6/latest/set-up/install-k6/','https://grafana.com/docs/k6/latest/results-output/end-of-test/custom-summary/','https://grafana.com/docs/k6/latest/results-output/real-time/json/','https://www.postgresql.org/download/windows/','https://www.postgresql.org/docs/16/transaction-iso.html','https://fastapi.tiangolo.com/async/']:p(link)
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output);print(args.output)

if __name__=='__main__':main()
